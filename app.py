import base64
import copy
import html
import io
import re
import zipfile
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

BASE = Path(__file__).parent
DATA_FILE = BASE / "Data.xlsx"
TEMPLATE = BASE / "Naveen_Job_Card_Empty.xlsx"

FLOORS = [
    ("First Floor Pricing", "FF"),
    ("First Floor with Basement Pricing", "FF+Bsmt"),
    ("Triplex (1st with basement and stilt room)", "Triplex"),
    ("Second Floor Pricing", "SF"),
    ("Third Floor Pricing", "TF"),
    ("Fourth Floor Pricing", "Top"),
    ("Fourth Floor with Terrace Pricing", "Top (Terrace)"),
    ("Duplex (4th with terrace)", "Duplex"),
]
CARD_COLS = ["Date", "Address", "Location", "Property", "Size", "Floor/s", "Price",
             "Facing", "Road", "Source Name", "Phone No.", "Type", "Tron Remarks",
             "Client Reaction"]
EXCEL_COLS = "ABCDEFGHIJKLMN"
RM_OPTIONS = ["Nikhil", "Tushar", "Rohil", "Pavi"]
FLAGS = [("Others [Corner]", "Corner"), ("Others [Park Facing]", "Park Facing"),
         ("Others [T-Point]", "T-Point"), ("Others [Resale]", "Resale"),
         ("Others [Gated]", "Gated")]


def parse_price(v):
    """Read a price cell. Returns (list of prices in Cr, display text) or None.

    Handles: 3.25 Cr, 4.2cr, 3.7, "3,50" (decimal comma), "10. 0 Cr",
    "4.45 Cr With half terrace" (note ignored), ranges like "4.50 Cr-5.0 Cr",
    and lists like "6.5, 6.8". Ignores things like "1/4 basement" and
    "sold", "hold", "GF-", "Full building" entries.
    """
    if v is None or pd.isna(v):
        return None
    if isinstance(v, (int, float)):
        return [float(v)], "%g" % v
    s = str(v).strip()
    if re.search(r"sold|hold|full|all building|gf|^g\b|^g-", s, re.I):
        return None
    s = re.sub(r"\d+\s*/\s*\d+", " ", s)             # drop fractions like 1/4
    s = re.sub(r"(\d)\.\s+(\d)", r"\1.\2", s)         # "10. 0" -> "10.0"
    if re.fullmatch(r"\s*\d+,\d{1,2}\s*(cr)?\s*", s, re.I):
        s = s.replace(",", ".")                          # "3,50" -> "3.50"
    nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", s)]
    nums = [n for n in nums if 0.5 <= n <= 500]
    if not nums:
        return None
    is_range = len(nums) > 1 and re.search(r"\d\s*(cr)?\s*(-|to)\s*\d", s, re.I)
    sep = "-" if is_range else "/"
    return nums, sep.join("%g" % n for n in nums)


def first_number(v):
    m = re.search(r"\d+", str(v))
    return int(m.group()) if m else 0


def to_ts(v):
    if isinstance(v, (int, float)) and not pd.isna(v):
        return pd.Timestamp("1899-12-30") + pd.to_timedelta(v, unit="D")
    if hasattr(v, "year"):
        return pd.Timestamp(v)
    return pd.to_datetime(str(v), dayfirst=True, errors="coerce")


def is_red(cell):
    """True if the cell has a red fill or red text (our 'sold' marker)."""
    def red_rgb(c):
        try:
            if c is None or c.type != "rgb" or not isinstance(c.rgb, str):
                return False
            h = c.rgb[-6:]
            r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
            return r >= 0xC0 and g <= 0x60 and b <= 0x60
        except Exception:
            return False
    if cell.fill is not None and cell.fill.fill_type == "solid" and red_rgb(cell.fill.fgColor):
        return True
    return cell.font is not None and red_rgb(cell.font.color)


SOLD_ROW_MIN_RED_CELLS = 15   # a row with this many red cells counts as "entire row red"


@st.cache_data
def load_data():
    df = pd.read_excel(DATA_FILE)
    df["_size"] = df["Plot Size (Sq Yard)"].map(first_number)
    df["_ts"] = df["Timestamp"].map(to_ts)

    # Read the red colouring straight from the Excel file (pandas can't see colours)
    ws = load_workbook(DATA_FILE).active
    header = [c.value for c in ws[1]]
    col_idx = {name: i + 1 for i, name in enumerate(header)}
    sold_row, red_floor = {}, {}
    for r in range(2, ws.max_row + 1):
        n_red = sum(1 for c in range(1, len(header) + 1) if is_red(ws.cell(r, c)))
        sold_row[r] = n_red >= SOLD_ROW_MIN_RED_CELLS
        red_floor[r] = {lab for col, lab in FLOORS
                        if col in col_idx and is_red(ws.cell(r, col_idx[col]))}

    df["_sold"] = [sold_row.get(i + 2, False) for i in df.index]

    def floors_of(r):
        out = []
        reds = red_floor.get(r.name + 2, set())
        for col, lab in FLOORS:
            p = parse_price(r[col])
            if p:
                out.append((lab, p[0], p[1], lab in reds))
        return out

    df["_floors"] = df.apply(floors_of, axis=1)
    return df


def clean(v):
    return "" if v is None or (not isinstance(v, str) and pd.isna(v)) else v


def make_row(r, floors, visit_date):
    rem = [str(r["Stage of Property"])]
    c = r["Other Info or Comments"]
    if isinstance(c, str) and c.strip():
        rem.append(c.strip())
    for col, lab in FLAGS:
        if str(r[col]).strip().lower() == "yes":
            rem.append(lab)
    seen, out = set(), []
    for x in rem:
        if x.lower() not in seen:
            seen.add(x.lower())
            out.append(x)
    out = [x for x in out
           if not (x in dict(FLAGS).values()
                   and any(x.lower() in y.lower() and y != x for y in out))]
    phone = r["Source Contact No."]
    if isinstance(phone, float) and not pd.isna(phone):
        phone = int(phone)
    road = r["Road Size"]
    return {
        "Date": visit_date,
        "Address": r["Plot No."],
        "Location": r["Locality"],
        "Property": str(r["Rooms"]).replace(" ", ""),
        "Size": r["Plot Size (Sq Yard)"],
        "Floor/s": ", ".join(f for f, _, _ in floors),
        "Price": ", ".join(t for _, _, t in floors),
        "Facing": clean(r["Plot Facing"]),
        "Road": int(road) if not pd.isna(road) else "",
        "Source Name": r["Source Name"],
        "Phone No.": clean(phone),
        "Type": r["Source of Building"],
        "Tron Remarks": "; ".join(out),
        "Client Reaction": "",
    }


TEAL = PatternFill("solid", fgColor="00897B")
ZEBRA = PatternFill("solid", fgColor="EAF5F3")
NO_FILL = PatternFill(fill_type=None)
SIDE = Side(style="thin", color="333333")
BOX = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
WHITE_BOLD = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Calibri", size=11)
WRAP_COLS = "CFJM"                       # columns whose text wraps onto a second line
# (min width, max width) for each column A..N
COL_LIMITS = {"A": (12, 14), "B": (11, 16), "C": (16, 26), "D": (10, 12), "E": (8, 10),
              "F": (14, 24), "G": (12, 22), "H": (8, 10), "I": (7, 9), "J": (14, 24),
              "K": (13, 15), "L": (10, 12), "M": (22, 36), "N": (22, 30)}


def build_card(rows, header):
    wb = load_workbook(TEMPLATE)
    ws = wb["Naveen"]
    ws["A4"] = f"Client Name:   {header['name']}"
    ws["E4"] = f"Budget Range: {header['budget']}"
    ws["I4"] = f"Client Type:  {header['ctype']}"
    ws["M4"] = f"RM: {header['rm']}"
    ws["A5"] = f"Contact Number: {header['contact']}"
    ws["E5"] = f"Preferred Micro-Market:  {header['markets']}"

    # Client band (rows 4-5): teal, white bold text, one consistent dark border all round
    for r in (4, 5):
        for col in EXCEL_COLS:
            cell = ws[f"{col}{r}"]
            cell.fill = TEAL
            cell.font = WHITE_BOLD
            cell.border = BOX
            cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r].height = 20

    # Row 6 is only a thin spacer: no boxes, no borders
    for col in EXCEL_COLS:
        cell = ws[f"{col}6"]
        cell.value = None
        cell.fill = NO_FILL
        cell.border = Border()
    ws.row_dimensions[6].height = 6

    # Column headings (row 7)
    for col in EXCEL_COLS:
        cell = ws[f"{col}7"]
        cell.fill = TEAL
        cell.font = WHITE_BOLD
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[7].height = 24

    # Data rows: identical font, border and alignment everywhere, light zebra stripes
    last = 7 + len(rows)
    for r in range(8, max(ws.max_row, last) + 1):
        ws.row_dimensions[r].hidden = False      # the old template had rows hidden
        ws.row_dimensions[r].height = None
    for i, row in enumerate(rows):
        r = 8 + i
        for col, key in zip(EXCEL_COLS, CARD_COLS):
            v = row.get(key, "")
            cell = ws[f"{col}{r}"]
            cell.value = v if v != "" else None
            cell.font = BODY_FONT
            cell.border = BOX
            cell.fill = ZEBRA if i % 2 else NO_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=col in WRAP_COLS)

    # Remove every spare empty row below the data
    if ws.max_row > last:
        ws.delete_rows(last + 1, ws.max_row - last)
    for r in [k for k in ws.row_dimensions if k > last]:
        del ws.row_dimensions[r]

    # Column widths that fit the content
    for k in list(ws.column_dimensions.keys()):
        del ws.column_dimensions[k]
    for col, key in zip(EXCEL_COLS, CARD_COLS):
        lo, hi = COL_LIMITS[col]
        longest = max([len(key)] + [len(str(row.get(key, ""))) for row in rows])
        ws.column_dimensions[col].width = max(lo, min(hi, longest + 3))

    # Print setup: A4 landscape, 0.25 inch margins, one page wide, headings repeat
    ws.sheet_view.showGridLines = False
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.25, bottom=0.25,
                                  header=0.1, footer=0.1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "7:7"
    ws.print_area = f"A1:N{last}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- UI
@st.cache_data
def load_logo():
    """Pull the Tron logo out of the job-card template (no extra file needed)."""
    try:
        with zipfile.ZipFile(TEMPLATE) as z:
            media = [n for n in z.namelist() if n.startswith("xl/media/")]
            return z.read(media[0]) if media else None
    except Exception:
        return None


logo_bytes = load_logo()
page_icon = "🏢"
if logo_bytes:
    try:
        from PIL import Image
        page_icon = Image.open(io.BytesIO(logo_bytes))
    except Exception:
        pass

st.set_page_config(page_title="Tron Job Card Builder", page_icon=page_icon, layout="wide")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
:root {
  --teal:#14B8A6; --teal-600:#0D9488; --teal-700:#0F766E; --teal-900:#064E4A;
  --bg:#0B1F1E; --card:#102A28; --card-2:#143633; --line:#1F4A46;
  --text:#E6F4F1; --muted:#8FB5B0;
}
html, body, .stApp, button, input, textarea { font-family:'Inter', system-ui, -apple-system, 'Segoe UI', sans-serif; }
.stApp { background: radial-gradient(1100px 520px at 8% -8%, #14504B 0%, rgba(11,31,30,0) 60%) fixed, var(--bg); color:var(--text); }
#MainMenu, footer { visibility:hidden; }
[data-testid="stHeader"] { background:transparent; }
.block-container { padding-top:1.4rem; padding-bottom:2rem; max-width:1360px; }

/* ---------- hero ---------- */
.hero { display:flex; align-items:center; gap:22px; padding:22px 28px; border-radius:16px;
        background:linear-gradient(120deg, var(--teal-900) 0%, var(--teal-700) 55%, var(--teal-600) 100%);
        border:1px solid rgba(255,255,255,.12); box-shadow:0 10px 30px rgba(0,0,0,.35); }
.hero .logo { height:74px; width:74px; object-fit:contain; background:#fff; border-radius:14px; padding:6px;
              box-shadow:0 4px 14px rgba(0,0,0,.3); }
.hero .eyebrow { font-size:.72rem; letter-spacing:.22em; font-weight:600; color:#B2F5EA; text-transform:uppercase; }
.hero h1 { margin:.15rem 0 .25rem 0; padding:0; font-size:1.95rem; font-weight:800; color:#fff; letter-spacing:-.02em; line-height:1.15; }
.hero p { margin:0; color:#D5F5F0; font-size:.95rem; font-weight:400; }
.hero .badge { margin-left:auto; text-align:right; font-size:.72rem; color:#B2F5EA; letter-spacing:.12em; text-transform:uppercase; font-weight:600; }

/* ---------- KPI cards ---------- */
.kpis { display:grid; grid-template-columns:repeat(4, 1fr); gap:14px; margin:18px 0 10px 0; }
.kpi { background:linear-gradient(180deg, var(--card-2), var(--card)); border:1px solid var(--line); border-radius:14px; padding:16px 18px; }
.kpi .k-label { font-size:.72rem; letter-spacing:.1em; text-transform:uppercase; color:var(--muted); font-weight:600; }
.kpi .k-value { font-size:1.75rem; font-weight:800; color:#fff; margin-top:4px; letter-spacing:-.02em; }
.kpi .k-value small { font-size:.8rem; font-weight:600; color:var(--teal); margin-left:4px; }
@media (max-width:900px){ .kpis { grid-template-columns:repeat(2,1fr); } .hero { flex-wrap:wrap; } .hero .badge{display:none;} }

/* ---------- section titles ---------- */
.sec-title { font-size:1.05rem; font-weight:700; color:#fff; letter-spacing:-.01em; margin:0; }
.sec-sub { font-size:.82rem; color:var(--muted); margin:2px 0 0 0; }
.mini-label { font-size:.78rem; font-weight:600; color:var(--muted); letter-spacing:.03em; text-transform:uppercase; margin:0 0 .45rem 0; }

/* ---------- bordered containers (filters / results) ---------- */
[data-testid="stVerticalBlockBorderWrapper"] { background:linear-gradient(180deg, var(--card), #0E2523);
        border:1px solid var(--line) !important; border-radius:16px; box-shadow:0 6px 22px rgba(0,0,0,.25); }

/* ---------- inputs ---------- */
[data-testid="stWidgetLabel"] p, label p { font-size:.78rem !important; font-weight:600 !important; color:var(--muted) !important;
        letter-spacing:.03em; text-transform:uppercase; }
[data-baseweb="select"] > div, [data-baseweb="input"], [data-baseweb="base-input"], .stNumberInput input, .stTextInput input {
        background:#0B1F1E !important; border-color:var(--line) !important; border-radius:10px !important; color:var(--text) !important; }
[data-baseweb="select"] > div:hover, [data-baseweb="input"]:hover { border-color:var(--teal-600) !important; }
[data-baseweb="tag"] { background:var(--teal-600) !important; color:#fff !important; border-radius:6px !important; font-weight:600; }
[data-baseweb="tag"] span { color:#fff !important; }
[data-testid="stMultiSelectTagsContainer"] > span > span { background:var(--teal-600) !important; color:#fff !important; border-radius:6px !important; font-weight:600; }
[data-testid="stMultiSelectTagsContainer"] > span > span * { color:#fff !important; }
[data-testid="stCheckbox"] label p { text-transform:none !important; letter-spacing:0; font-size:.88rem !important; color:var(--text) !important; font-weight:500 !important; }

/* ---------- buttons ---------- */
.stButton > button { background:transparent !important; color:#5EEAD4 !important; border:1px solid var(--teal-600) !important;
        border-radius:10px; font-weight:600; padding:.45rem 1rem; transition:all .15s ease; }
.stButton > button p { color:#5EEAD4 !important; }
.stButton > button:hover { background:rgba(20,184,166,.12) !important; border-color:var(--teal) !important; }
.stDownloadButton > button { background:linear-gradient(120deg, var(--teal-700), var(--teal)) !important; color:#fff !important; border:0 !important;
        border-radius:12px; font-weight:700; padding:.65rem 1.4rem; box-shadow:0 6px 18px rgba(20,184,166,.35); transition:all .15s ease; }
.stDownloadButton > button p { color:#fff !important; font-weight:700; }
.stDownloadButton > button:hover { transform:translateY(-1px); box-shadow:0 10px 24px rgba(20,184,166,.5); }

/* ---------- sidebar ---------- */
[data-testid="stSidebar"] { background:linear-gradient(190deg, #0B3F3A 0%, #082E2B 100%); border-right:1px solid var(--line); }
[data-testid="stSidebar"] h2 { color:#fff !important; font-size:1.05rem; font-weight:700; letter-spacing:.02em; }
[data-testid="stSidebar"] label p { color:#9ED9D1 !important; }
[data-testid="stSidebar"] [data-baseweb="input"], [data-testid="stSidebar"] [data-baseweb="base-input"],
[data-testid="stSidebar"] [data-testid="stTextInputRootElement"] { background:#F4FBFA !important; border:1px solid #9ED9D1 !important; border-radius:10px; }
[data-testid="stSidebar"] [data-testid="stMultiSelect"] > div > div { background:#F4FBFA !important; border:1px solid #9ED9D1 !important; border-radius:10px; }
[data-testid="stSidebar"] [data-testid="stMultiSelect"] svg { color:#0B1F1E !important; fill:#0B1F1E !important; }
[data-testid="stSidebar"] input { color:#0B1F1E !important; -webkit-text-fill-color:#0B1F1E !important; caret-color:#0B1F1E; background:transparent !important; font-weight:500; }
[data-testid="stSidebar"] input::placeholder { color:#6B8F8A !important; -webkit-text-fill-color:#6B8F8A !important; }
.side-note { font-size:.78rem; color:#7FC4BB; line-height:1.5; margin-top:6px; }

/* ---------- table ---------- */
[data-testid="stDataFrame"] { border:1px solid var(--line); border-radius:12px; overflow:hidden; }

/* ---------- alerts ---------- */
[data-testid="stAlert"] { background:rgba(20,184,166,.10) !important; border:1px solid var(--line); border-radius:12px; }
[data-testid="stAlert"] * { color:#CFF5EF !important; }

/* ---------- footer ---------- */
.app-footer { text-align:center; color:var(--muted); font-size:.76rem; letter-spacing:.06em; margin-top:26px; padding-top:14px; border-top:1px solid var(--line); }
</style>
""", unsafe_allow_html=True)

df = load_data()

# ---------------------------------------------------------------- sidebar
with st.sidebar:
    st.header("Client details")
    name = st.text_input("Client name")
    budget = st.text_input("Budget range", placeholder="e.g. 7.5 Cr")
    ctype = st.text_input("Client type")
    rm = st.multiselect("RM", RM_OPTIONS, placeholder="Select RM")
    contact = st.text_input("Contact number")
    visit_date = st.text_input("Visit date (optional)", placeholder="DD/MM/YYYY")
    st.markdown('<div class="side-note">These details are printed in the header of the '
                'downloaded job card.</div>', unsafe_allow_html=True)

# ---------------------------------------------------------------- hero
logo_tag = ""
if logo_bytes:
    logo_tag = ('<img class="logo" alt="Tron Asset Management" src="data:image/png;base64,'
                + base64.b64encode(logo_bytes).decode() + '">')
st.markdown(f"""
<div class="hero">
  {logo_tag}
  <div>
    <div class="eyebrow">Tron Asset Management LLP</div>
    <h1>Client Site Visit Job Card Builder</h1>
    <p>Search the property database, shortlist the right listings and generate a branded job card in seconds.</p>
  </div>
  <div class="badge">Internal<br>Advisory Tool</div>
</div>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------- filters
INITIAL = {"loc": [], "addr": [], "facing": [], "rooms": ["4 BHK"],
           "stage": ["RTM", "U/C"], "sources": [], "road": [], "others": [], "floor_pref": [], "lift_stilt": True, "min_size": 0, "max_size": 0,
           "min_price": 0.0, "max_price": 0.0, "latest_only": True,
           "hide_rows": True, "hide_floors": True}
CLEARED = {**INITIAL, "rooms": [], "stage": []}
for k, v in INITIAL.items():
    st.session_state.setdefault(k, v)


def clear_filters():
    for k, v in CLEARED.items():
        st.session_state[k] = v


def norm(x):
    return re.sub(r"[^a-z0-9]", "", str(x).lower())


@st.cache_data
def address_options(frame):
    """Labels like 'A-1084 · Sushant Lok 1' -> (normalised plot no., locality)."""
    pairs = frame[["Plot No.", "Locality"]].dropna().astype(str).drop_duplicates()
    lookup = {f"{p.strip()} · {l}": (norm(p), l) for p, l in zip(pairs["Plot No."], pairs["Locality"])}
    return sorted(lookup), lookup


addr_labels, addr_lookup = address_options(df[["Plot No.", "Locality"]])

OTHERS = ["Corner", "Resale", "Gated", "T-Point", "Park Facing"]
FLOOR_NAMES = {"FF": "First Floor (FF)", "FF+Bsmt": "First Floor + Basement", "Triplex": "Triplex",
               "SF": "Second Floor (SF)", "TF": "Third Floor (TF)", "Top": "Fourth / Top Floor",
               "Top (Terrace)": "Top Floor with Terrace", "Duplex": "Duplex"}
road_options = sorted({int(x) for x in df["Road Size"].dropna()})

kpi_slot = st.container()

with st.container(border=True):
    t1, t2 = st.columns([6, 1])
    t1.markdown('<p class="sec-title">Filters</p><p class="sec-sub">Narrow the database by area, '
                'size, price, facing and more.</p>', unsafe_allow_html=True)
    t2.button("Clear all", on_click=clear_filters, width="stretch")

    c1, c2, c3 = st.columns(3)
    with c1:
        localities = st.multiselect("Location / area", sorted(df["Locality"].unique()), key="loc")
        addr = st.multiselect("Address / plot no. / name", addr_labels, key="addr",
                              accept_new_options=True,
                              placeholder="Pick or type, press Enter to add more")
        facing = st.multiselect("Facing", sorted(df["Plot Facing"].dropna().unique()), key="facing")
    with c2:
        rooms = st.multiselect("Property type", sorted(df["Rooms"].dropna().unique()), key="rooms")
        stage = st.multiselect("Stage", ["RTM", "U/C"], key="stage")
        sources = st.multiselect("Source type", sorted(df["Source of Building"].unique()),
                                 key="sources")
    with c3:
        s1, s2 = st.columns(2)
        min_size = s1.number_input("Min size (yd)", min_value=0, step=10, key="min_size",
                                   help="Plot size in square yards")
        max_size = s2.number_input("Max size (yd)", min_value=0, step=10, key="max_size",
                                   help="Plot size in square yards. 0 means no limit.")
        p1, p2 = st.columns(2)
        min_price = p1.number_input("Min price (Cr)", min_value=0.0, step=0.25, key="min_price")
        max_price = p2.number_input("Max price (Cr)", min_value=0.0, step=0.25,
                                    key="max_price", help="0 means no limit.")
        road = st.multiselect("Road size", road_options, key="road",
                              format_func=lambda x: f"{x:g}")

    f1, f2, f3 = st.columns(3)
    floor_pref = f1.multiselect("Floor preference", list(FLOOR_NAMES), key="floor_pref",
                                format_func=lambda x: FLOOR_NAMES[x],
                                placeholder="Any floor",
                                help="Shows only the selected floors, and only properties that have one of them available.")
    others = f2.multiselect("Others (property features)", OTHERS, key="others",
                            placeholder="Pick features",
                            help="Shows only properties that have ALL the selected features.")
    f3.markdown('<div class="mini-label">Required features</div>', unsafe_allow_html=True)
    lift_stilt = f3.checkbox("Lift & Stilt Parking required", key="lift_stilt")

    o1, o2, o3 = st.columns(3)
    latest_only = o1.checkbox("Latest entry per plot only", key="latest_only")
    hide_sold_rows = o2.checkbox("Hide sold properties (entire row red)", key="hide_rows")
    hide_sold_floors = o3.checkbox("Hide sold floors (red price cell)", key="hide_floors")

# ---------------------------------------------------------------- filtering
d = df.copy()
if latest_only:
    d = d.sort_values("_ts", ascending=False).drop_duplicates(["Plot No.", "Locality"])
n_sold_hidden = 0
if hide_sold_rows:
    n_before = len(d)
    d = d[~d["_sold"]]
    n_sold_hidden = n_before - len(d)
if localities:
    d = d[d["Locality"].isin(localities)]
if facing:
    d = d[d["Plot Facing"].isin(facing)]
if rooms:
    d = d[d["Rooms"].isin(rooms)]
if stage:
    d = d[d["Stage of Property"].isin(stage)]
if sources:
    d = d[d["Source of Building"].isin(sources)]
if road:
    d = d[d["Road Size"].isin(road)]
if lift_stilt:
    for col in ("Others [Lift]", "Others [Stilt Parking]"):
        d = d[d[col].astype(str).str.strip().str.lower() == "yes"]
for feat in others:
    d = d[d[f"Others [{feat}]"].astype(str).str.strip().str.lower() == "yes"]
if min_size:
    d = d[d["_size"] >= min_size]
if max_size:
    d = d[d["_size"] <= max_size]
if addr:
    picked = {addr_lookup[a] for a in addr if a in addr_lookup}      # exact plot + area
    typed = [a.strip().lower() for a in addr if a not in addr_lookup and a.strip()]
    keep = pd.Series(False, index=d.index)
    if picked:
        keep |= pd.Series([(norm(p), l) in picked for p, l in
                           zip(d["Plot No."].astype(str), d["Locality"].astype(str))],
                          index=d.index)
    if typed:                                                         # free text: search
        hay = (d["Plot No."].astype(str) + " " + d["Locality"].astype(str) + " "
               + d["Source Name"].astype(str) + " " + d["Blocks"].astype(str)).str.lower()
        for t in typed:
            keep |= hay.str.contains(re.escape(t))
    d = d[keep]

rows, all_prices = [], []
for _, r in d.iterrows():
    fl = [(lab, nums, txt) for lab, nums, txt, red in r["_floors"]
          if not (hide_sold_floors and red and not r["_sold"])
          and (not floor_pref or lab in floor_pref)
          and any(n >= min_price and (not max_price or n <= max_price) for n in nums)]
    if not fl:
        continue
    all_prices += [n for _, nums, _ in fl for n in nums]
    rows.append(make_row(r, fl, visit_date))

rows.sort(key=lambda x: (x["Location"], first_number(x["Size"]), str(x["Address"])))

# ---------------------------------------------------------------- KPI cards
areas = len({x["Location"] for x in rows})
price_rng = (f"{min(all_prices):g} – {max(all_prices):g}<small>Cr</small>" if all_prices else "–")
with kpi_slot:
    st.markdown(f"""
<div class="kpis">
  <div class="kpi"><div class="k-label">Matching properties</div><div class="k-value">{len(rows)}</div></div>
  <div class="kpi"><div class="k-label">Areas covered</div><div class="k-value">{areas}</div></div>
  <div class="kpi"><div class="k-label">Listed price range</div><div class="k-value">{price_rng}</div></div>
  <div class="kpi"><div class="k-label">Sold listings hidden</div><div class="k-value">{n_sold_hidden if hide_sold_rows else 0}</div></div>
</div>
""", unsafe_allow_html=True)

# ---------------------------------------------------------------- results
st.write("")
with st.container(border=True):
    st.markdown('<p class="sec-title">Results</p><p class="sec-sub">Untick any property you do not '
                'want on the job card.</p>', unsafe_allow_html=True)
    if rows:
        view = pd.DataFrame(rows)
        for c in ("Size", "Road", "Phone No.", "Address"):
            view[c] = view[c].astype(str)
        # hide columns that are still blank (they are still filled in the Excel card)
        view = view.drop(columns=[c for c in ("Date", "Client Reaction")
                                  if (view[c].astype(str).str.strip() == "").all()])
        view.insert(0, "Select", True)
        edited = st.data_editor(
            view, hide_index=True, width="stretch", height=470,
            disabled=[c for c in view.columns if c != "Select"],
            column_config={"Select": st.column_config.CheckboxColumn("Include", width="small")},
        )
        chosen = [rows[i] for i in edited.index[edited["Select"]]]
        a1, a2 = st.columns([3, 2])
        a1.markdown(f'<p class="sec-sub" style="margin-top:14px">'
                    f'<b style="color:#fff">{len(chosen)}</b> of {len(rows)} properties selected '
                    f'for the job card</p>', unsafe_allow_html=True)
        if chosen:
            a2.download_button(
                "Download job card (Excel)",
                data=build_card(chosen, dict(name=name, budget=budget, ctype=ctype, rm=", ".join(rm),
                                             contact=contact,
                                             markets=", ".join(localities))),
                file_name=f"Job_Card_{(name or 'Client').strip().replace(' ', '_')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                width="stretch",
            )
    else:
        st.info("No properties match these filters. Try widening the price or size range, "
                "or press Clear all.")

st.markdown('<div class="app-footer">TRON ASSET MANAGEMENT LLP &nbsp;·&nbsp; Confidential '
            '&nbsp;·&nbsp; For internal advisory use only</div>', unsafe_allow_html=True)
